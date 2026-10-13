
from pathlib import Path
from datetime import datetime
import tempfile
import re

import joblib
import pandas as pd
import pytesseract

from PIL import Image, ImageOps, ImageEnhance
from sklearn.metrics import accuracy_score

# ==================================================
# 1. PROJECT SETTINGS
# ==================================================

BASE_DIR = Path(__file__).resolve().parent
DATASET_DIR = BASE_DIR / "dataset"
HARD_TEST_DIR = DATASET_DIR / "hard_test"
MODEL_DIR = BASE_DIR / "models"

MODEL_PATH = MODEL_DIR / "xgboost_document_model.joblib"
VECTORIZER_PATH = MODEL_DIR / "tfidf_vectorizer.joblib"
LABELS_PATH = MODEL_DIR / "document_labels.joblib"

# Save outside OneDrive to avoid overwriting locked reports.
OUTPUT_DIR = Path(tempfile.gettempdir())

pytesseract.pytesseract.tesseract_cmd = (
    r"C:\Program Files\Tesseract-OCR\tesseract.exe"
)

SUPPORTED_EXTENSIONS = {
    ".png", ".jpg", ".jpeg",
    ".bmp", ".tif", ".tiff"
}

# Actual labels for hard_01.png through hard_20.png.
# Confirm these labels match your hard_test images.
EXPECTED_LABELS = {
    "hard_01": "aadhar",
    "hard_02": "pan",
    "hard_03": "driving_license",
    "hard_04": "unknown",
    "hard_05": "aadhar",
    "hard_06": "pan",
    "hard_07": "driving_license",
    "hard_08": "unknown",
    "hard_09": "aadhar",
    "hard_10": "pan",
    "hard_11": "driving_license",
    "hard_12": "unknown",
    "hard_13": "aadhar",
    "hard_14": "pan",
    "hard_15": "driving_license",
    "hard_16": "unknown",
    "hard_17": "aadhar",
    "hard_18": "pan",
    "hard_19": "driving_license",
    "hard_20": "unknown",
}


# ==================================================
# 2. OCR AND TEXT CLEANING
# ==================================================

def clean_text(text):
    if not text:
        return ""

    return " ".join(text.split()).strip().lower()


def extract_text(image_path):
    with Image.open(image_path) as source:
        image = ImageOps.exif_transpose(source)
        image = image.convert("RGB")

        gray = ImageOps.grayscale(image)
        gray = ImageEnhance.Contrast(gray).enhance(1.5)

        text = pytesseract.image_to_string(
            gray,
            config="--psm 6"
        )

    return clean_text(text)


# ==================================================
# 3. LOAD TRAINED MODEL
# ==================================================

def load_model_files():
    for path in [MODEL_PATH, VECTORIZER_PATH, LABELS_PATH]:
        if not path.is_file():
            raise FileNotFoundError(
                f"Required model file not found: {path}"
            )

    model = joblib.load(MODEL_PATH)
    vectorizer = joblib.load(VECTORIZER_PATH)
    labels = joblib.load(LABELS_PATH)

    print("XGBoost model loaded successfully.")
    print("TF-IDF vectorizer loaded successfully.")
    print("Document labels loaded successfully.")

    return model, vectorizer, labels


# ==================================================
# 4. CONVERT MODEL OUTPUT TO LABEL
# ==================================================

def get_label(class_id, labels):
    # Training script saves ID-to-label dictionary.
    if isinstance(labels, dict):
        if class_id in labels:
            return labels[class_id]

        if str(class_id) in labels:
            return labels[str(class_id)]

    # Also support a list of labels, if used.
    if isinstance(labels, (list, tuple)):
        return labels[int(class_id)]

    raise ValueError(
        f"Cannot map model class ID {class_id} to a label."
    )


# ==================================================
# 5. PROCESS TEST IMAGES
# ==================================================

def generate_report():
    print("\n========== DOCUMENT CLASSIFICATION REPORT ==========\n")

    if not HARD_TEST_DIR.is_dir():
        raise FileNotFoundError(
            f"Hard-test folder not found: {HARD_TEST_DIR}"
        )

    model, vectorizer, labels = load_model_files()

    image_paths = sorted(
        path for path in HARD_TEST_DIR.iterdir()
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )

    if not image_paths:
        raise ValueError(
            f"No supported images found in {HARD_TEST_DIR}"
        )

    report_rows = []
    actual_labels = []
    predicted_labels = []

    for index, image_path in enumerate(image_paths, start=1):
        print(
            f"Processing {index}/{len(image_paths)}: "
            f"{image_path.name}"
        )

        try:
            text = extract_text(image_path)

            if not text:
                print("  Warning: No OCR text detected.")

            features = vectorizer.transform([text])

            predicted_id = int(model.predict(features)[0])
            predicted_label = get_label(predicted_id, labels)

            probabilities = model.predict_proba(features)[0]
            confidence = float(max(probabilities)) * 100

            # Get the true label from the known test mapping.
            # Unknown filenames must not be guessed.
            actual_label = EXPECTED_LABELS.get(
                image_path.stem.lower()
            )

            if actual_label is None:
                match_status = "LABEL REQUIRED"
            else:
                match_status = (
                    "MATCH"
                    if actual_label == predicted_label
                    else "NOT MATCHED"
                )

                actual_labels.append(actual_label)
                predicted_labels.append(predicted_label)

            report_rows.append({
                "SI No": len(report_rows) + 1,
                "Document Name": image_path.name,
                "Document Type": actual_label or "LABEL REQUIRED",
                "Model Prediction": predicted_label,
                "Confidence": round(confidence, 2),
                "Accuracy": match_status,
            })

        except Exception as error:
            print(f"  ERROR: {error}")

            actual_label = EXPECTED_LABELS.get(
                image_path.stem.lower(),
                "LABEL REQUIRED"
            )

            report_rows.append({
                "SI No": len(report_rows) + 1,
                "Document Name": image_path.name,
                "Document Type": actual_label,
                "Model Prediction": "ERROR",
                "Confidence": 0.0,
                "Accuracy": "NOT EVALUATED",
            })

    if not report_rows:
        raise RuntimeError("No report rows could be generated.")

    # ==================================================
    # 6. CREATE EXCEL REPORT
    # ==================================================

    df = pd.DataFrame(report_rows)

    # Unique filename prevents overwriting a locked report.
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = (
        OUTPUT_DIR
        / f"document_classification_report_{timestamp}.xlsx"
    )

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Classification Report",
                    index=False)

        worksheet = writer.sheets["Classification Report"]

        # Set column widths.
        widths = {
            "A": 10,
            "B": 24,
            "C": 22,
            "D": 24,
            "E": 16,
            "F": 20,
        }

        for column, width in widths.items():
            worksheet.column_dimensions[column].width = width

        # Keep the header visible while scrolling.
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        # Format confidence as a percentage value.
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row=row, column=5).number_format = '0.00"%"'

        # Highlight header.
        from openpyxl.styles import Font, PatternFill, Alignment

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="1F4E78"
            )
            cell.alignment = Alignment(horizontal="center")

    # ==================================================
    # 7. DISPLAY ACCURACY
    # ==================================================

    evaluated_count = len(actual_labels)
    correct_count = sum(
        actual == predicted
        for actual, predicted in zip(
            actual_labels, predicted_labels
        )
    )

    print("\n========== REPORT COMPLETED ==========")
    print(f"Documents processed: {len(report_rows)}")
    print(f"Documents evaluated: {evaluated_count}")
    print(f"Correct predictions: {correct_count}")
    print(f"Incorrect predictions: {evaluated_count - correct_count}")

    if evaluated_count:
        accuracy = accuracy_score(actual_labels, predicted_labels)
        print(f"Overall test accuracy: {accuracy:.2%}")
    else:
        print("Overall test accuracy: Not available")

    print(f"\nExcel report saved successfully:\n{output_path}")
    print("\nOpen the report using the path shown above.")


# ==================================================
# 8. MAIN
# ==================================================

if __name__ == "__main__":
    try:
        generate_report()
    except PermissionError as error:
        print("\nPERMISSION ERROR: Windows blocked file access.")
        print("Close Excel and check Windows/OneDrive permissions.")
        print(error)
    except Exception as error:
        print(f"\nERROR: {error}")

